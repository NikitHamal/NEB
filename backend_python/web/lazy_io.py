import io
import re

from django.http import JsonResponse, HttpResponse


def _esc(s):
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


def strip_tags(html):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html or "")).strip()


def extract_docx(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    title = ""
    html_parts = []
    for para in doc.paragraphs:
        txt = (para.text or "").strip()
        if not txt:
            continue
        style = (para.style.name or "").lower()
        if not title and txt:
            title = txt[:120]
        if "heading 1" in style:
            html_parts.append(f"<h1>{_esc(txt)}</h1>")
        elif "heading 2" in style:
            html_parts.append(f"<h2>{_esc(txt)}</h2>")
        elif "heading 3" in style:
            html_parts.append(f"<h3>{_esc(txt)}</h3>")
        else:
            bold = any(r.bold for r in para.runs if r.text)
            if bold and len(txt) < 80:
                html_parts.append(f"<h2>{_esc(txt)}</h2>")
            else:
                html_parts.append(f"<p>{_esc(txt)}</p>")
    for table in getattr(doc, "tables", []):
        rows = [[_esc(c.text.strip()) for c in row.cells] for row in table.rows]
        if rows:
            html_parts.append("<table><thead><tr>" + "".join(f"<th>{c}</th>" for c in rows[0]) + "</tr></thead><tbody>")
            for r in rows[1:]:
                html_parts.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
            html_parts.append("</tbody></table>")
    html = "".join(html_parts)
    if not html:
        txts = [_esc(p.text.strip()) for p in doc.paragraphs if p.text.strip()]
        html = "".join(f"<p>{t}</p>" for t in txts[:120])
    return html, title


def extract_pdf_text(data):
    text = ""
    title = ""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages[:40]:
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            if t.strip():
                pages.append(t.strip())
        text = "\n\n".join(pages)
        if reader.metadata and getattr(reader.metadata, "title", None):
            title = str(reader.metadata.title or "")[:120]
    except Exception:
        pass
    if not text.strip():
        try:
            import pikepdf
            pdf = pikepdf.open(io.BytesIO(data))
            pages = []
            for p in list(pdf.pages)[:40]:
                try:
                    t = p.extract_text() if hasattr(p, "extract_text") else ""
                except Exception:
                    t = ""
                if t:
                    pages.append(t)
            text = "\n\n".join(pages)
        except Exception:
            pass
    return text, title


def text_to_html(text):
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    parts = []
    for i, para in enumerate(paras[:80]):
        para = re.sub(r"\s+", " ", para).strip()
        if not para:
            continue
        if i == 0 and len(para) < 100:
            parts.append(f"<h1>{_esc(para)}</h1>")
        elif len(para) < 70 and para.isupper():
            parts.append(f"<h2>{_esc(para.title())}</h2>")
        elif len(para) < 90 and re.match(r"^(chapter|section|part)\b", para, re.I):
            parts.append(f"<h2>{_esc(para)}</h2>")
        else:
            parts.append(f"<p>{_esc(para)}</p>")
    return "".join(parts)


def export_docx(html, title, safe):
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except Exception as exc:
        return JsonResponse({"error": f"docx export unavailable: {exc}"}, status=500)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    blocks = html_blocks(html)
    if not blocks:
        doc.add_paragraph(title)
    else:
        for tag, txt in blocks:
            if tag == "h1":
                p = doc.add_heading(txt, level=1)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(20)
                    run.font.color.rgb = RGBColor(0x0B, 0x1C, 0x30)
            elif tag == "h2":
                p = doc.add_heading(txt, level=2)
                for run in p.runs:
                    run.font.size = Pt(14)
                    run.font.color.rgb = RGBColor(0x00, 0x4A, 0xC6)
            elif tag == "h3":
                p = doc.add_heading(txt, level=3)
                for run in p.runs:
                    run.font.size = Pt(12)
            elif tag == "blockquote":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                run = p.add_run(txt)
                run.italic = True
                run.font.color.rgb = RGBColor(0x43, 0x4E, 0x62)
            elif tag == "li":
                doc.add_paragraph(txt, style="List Bullet")
            else:
                doc.add_paragraph(txt)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    resp["Content-Disposition"] = f'attachment; filename="{safe}.docx"'
    return resp


def export_pdf(html, title, safe):
    try:
        plain = strip_tags(html)
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    except Exception as exc:
        return JsonResponse({"error": f"pdf export unavailable: {exc}"}, status=500)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18*mm, bottomMargin=18*mm, leftMargin=18*mm, rightMargin=18*mm, title=title)
    styles = getSampleStyleSheet()
    title_s = ParagraphStyle("LzTitle", parent=styles["Title"], fontSize=18, leading=22, alignment=TA_CENTER, textColor="#0B1C30")
    h2_s = ParagraphStyle("LzH2", parent=styles["Heading2"], fontSize=13, leading=16, textColor="#004AC6", spaceBefore=10, spaceAfter=4)
    body_s = ParagraphStyle("LzBody", parent=styles["BodyText"], fontSize=10.5, leading=15, alignment=TA_JUSTIFY, spaceAfter=6)
    story = []
    blocks = html_blocks(html)
    if not blocks:
        story.append(Paragraph(title, title_s))
        story.append(Spacer(1, 8))
        story.append(Paragraph(plain[:2000] or "Empty document", body_s))
    else:
        for tag, txt in blocks:
            txt = txt.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if tag == "h1":
                story.append(Paragraph(txt, title_s))
                story.append(Spacer(1, 6))
            elif tag == "h2":
                story.append(Paragraph(txt, h2_s))
            elif tag == "h3":
                s3 = ParagraphStyle("LzH3", parent=h2_s, fontSize=11, textColor="#1E293B")
                story.append(Paragraph(txt, s3))
            elif tag == "blockquote":
                bq = ParagraphStyle("LzBQ", parent=body_s, leftIndent=12, textColor="#434E62", fontName="Helvetica-Oblique")
                story.append(Paragraph(txt, bq))
            else:
                story.append(Paragraph(txt, body_s))
    try:
        doc.build(story)
    except Exception as exc:
        return JsonResponse({"error": f"pdf build failed: {exc}"}, status=500)
    buf.seek(0)
    resp = HttpResponse(buf.getvalue(), content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{safe}.pdf"'
    return resp


def html_blocks(html):
    out = []
    pattern = re.compile(r'<(h1|h2|h3|p|li|blockquote)[^>]*>(.*?)</\1>', re.I | re.S)
    for m in pattern.finditer(html or ""):
        inner = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", m.group(2))).strip()
        if inner:
            out.append((m.group(1).lower(), inner))
    return out


def fix_html(html):
    if not html:
        return html
    out = html
    for a, b in ((" ,", ","), (" .", "."), (" :", ":"), (" ;", ";"), ("Fig :", "Fig:"), ("Table :", "Table:"), ("  ", " ")):
        out = out.replace(a, b)
    out = re.sub(r"<p>\s*</p>", "", out)
    out = re.sub(r"<h2>\s*</h2>", "", out)
    return out.strip()
