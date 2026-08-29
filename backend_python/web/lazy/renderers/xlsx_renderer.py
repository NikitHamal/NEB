"""Styled XLSX workbook renderer.

Produces formatted spreadsheets: bold frozen header rows with fills, bordered
cells, sensible column widths, numeric alignment, optional summary block and
native Excel charts (bar / line / pie) rather than static dumps.
"""

import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ACCENT = '1F4E5F'
ACCENT_SOFT = 'EAF1F4'
BORDER_COLOR = 'C9D2D8'

HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)
HEADER_FILL = PatternFill('solid', fgColor=ACCENT)
TITLE_FONT = Font(bold=True, size=14, color=ACCENT)
NOTE_FONT = Font(italic=True, size=9, color='5A6470')
TOTAL_FONT = Font(bold=True, size=11)

THIN = Side(style='thin', color=BORDER_COLOR)
BOX_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _clean(value, limit=None):
    text = re.sub(r'\s+', ' ', str(value or '')).strip()
    return text[:limit] if limit else text


def _as_row(value, width=None):
    if isinstance(value, (list, tuple)):
        row = list(value)
    elif isinstance(value, dict):
        row = list(value.values())
    else:
        row = [value]
    return [r if isinstance(r, (int, float)) else _clean(r, 500) for r in row][:width or 40]


def _coerce_number(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value or '').strip().replace(',', '')
    if not text:
        return None
    try:
        if text.endswith('%'):
            return float(text[:-1]) / 100
        return float(text) if '.' in text else int(text)
    except ValueError:
        return None


def normalize_workbook(raw, default_title='Workbook'):
    raw = raw if isinstance(raw, dict) else {}
    sheets = []
    for entry in (raw.get('sheets') or [])[:20]:
        if not isinstance(entry, dict):
            continue
        header = [_clean(h, 120) for h in (entry.get('header') or entry.get('columns') or [])][:40]
        rows = []
        for row in (entry.get('rows') or entry.get('data') or [])[:2000]:
            rows.append(_as_row(row, len(header) if header else None))
        chart = entry.get('chart') if isinstance(entry.get('chart'), dict) else None
        sheets.append({
            'name': _clean(entry.get('name') or f'Sheet{len(sheets) + 1}', 31) or f'Sheet{len(sheets) + 1}',
            'title': _clean(entry.get('title') or '', 160),
            'notes': _clean(entry.get('notes') or '', 400),
            'header': header,
            'rows': rows,
            'summary': [_as_row(r) for r in (entry.get('summary') or [])][:20],
            'chart': chart,
        })

    if not sheets:
        sheets = [{
            'name': 'Sheet1',
            'title': _clean(raw.get('title') or default_title, 160),
            'notes': '',
            'header': [_clean(h, 120) for h in (raw.get('header') or [])],
            'rows': [_as_row(r) for r in (raw.get('rows') or [])],
            'summary': [],
            'chart': None,
        }]

    return {'title': _clean(raw.get('title') or default_title, 160), 'sheets': sheets}


def render(spec, safe_name='workbook'):
    book = normalize_workbook(spec)
    wb = Workbook()
    wb.remove(wb.active)

    for index, sheet_spec in enumerate(book['sheets']):
        ws = wb.create_sheet(title=sheet_spec['name'][:31] or f'Sheet{index + 1}')
        _render_sheet(ws, sheet_spec)

    if not wb.sheetnames:
        wb.create_sheet('Sheet1')

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _render_sheet(ws, sheet_spec):
    header = sheet_spec.get('header') or []
    rows = sheet_spec.get('rows') or []
    row_cursor = 1

    if sheet_spec.get('title'):
        cell = ws.cell(row=row_cursor, column=1, value=sheet_spec['title'])
        cell.font = TITLE_FONT
        row_cursor += 1

    if sheet_spec.get('notes'):
        cell = ws.cell(row=row_cursor, column=1, value=sheet_spec['notes'])
        cell.font = NOTE_FONT
        row_cursor += 1

    if sheet_spec.get('title') or sheet_spec.get('notes'):
        row_cursor += 1

    header_row = row_cursor
    if header:
        for col_idx, name in enumerate(header, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=name)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BOX_BORDER
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[header_row].height = 22
        row_cursor += 1

    data_start = row_cursor
    for row in rows:
        for col_idx, value in enumerate(row, start=1):
            number = _coerce_number(value)
            cell = ws.cell(row=row_cursor, column=col_idx, value=number if number is not None else value)
            cell.border = BOX_BORDER
            if number is not None:
                cell.alignment = Alignment(horizontal='right')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        row_cursor += 1

    data_end = row_cursor - 1

    for entry in sheet_spec.get('summary') or []:
        for col_idx, value in enumerate(entry, start=1):
            number = _coerce_number(value)
            cell = ws.cell(row=row_cursor, column=col_idx, value=number if number is not None else value)
            cell.font = TOTAL_FONT
            cell.fill = PatternFill('solid', fgColor=ACCENT_SOFT)
            cell.border = BOX_BORDER
            if number is not None:
                cell.alignment = Alignment(horizontal='right')
        row_cursor += 1

    _autosize(ws, header, rows)
    if header and data_end >= data_start:
        ws.freeze_panes = ws.cell(row=data_start, column=1)
        ws.auto_filter.ref = f'A{header_row}:{get_column_letter(max(1, len(header)))}{data_end}'

    _add_chart(ws, sheet_spec, header_row, data_start, data_end)


def _autosize(ws, header, rows):
    column_count = max([len(header)] + [len(r) for r in rows] + [1])
    for col_idx in range(1, min(column_count, 40) + 1):
        longest = 0
        for row in ([header] if header else []) + rows[:200]:
            if col_idx - 1 < len(row):
                longest = max(longest, len(str(row[col_idx - 1] or '')))
        width = min(48, max(11, longest + 3))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _add_chart(ws, sheet_spec, header_row, data_start, data_end):
    chart_spec = sheet_spec.get('chart')
    if not chart_spec or data_end < data_start:
        return

    kind = str(chart_spec.get('type') or 'bar').strip().lower()
    title = _clean(chart_spec.get('title') or '', 120)
    col_count = max([len(sheet_spec.get('header') or [])] + [len(r) for r in sheet_spec.get('rows') or []] + [1])

    if kind == 'pie':
        chart = PieChart()
    elif kind == 'line':
        chart = LineChart()
    else:
        chart = BarChart()
        chart.type = 'col'

    chart.title = title or None
    chart.height = 8
    chart.width = 16

    if kind == 'pie':
        categories = Reference(ws, min_col=1, min_row=data_start, max_row=data_end)
        values = Reference(ws, min_col=min(2, col_count), min_row=data_start, max_row=data_end)
        chart.add_data(values, titles_from_data=False)
        chart.set_categories(categories)
    else:
        values = Reference(ws, min_col=2, max_col=col_count, min_row=header_row, max_row=data_end)
        categories = Reference(ws, min_col=1, min_row=data_start, max_row=data_end)
        chart.add_data(values, titles_from_data=True)
        chart.set_categories(categories)

    ws.add_chart(chart, f'A{data_end + 3}')
