"""Live end-to-end check: drive a real run against the real model.

Unlike `_e2e_check.py` (which stubs the LLM) this one costs credits and makes
genuine provider calls. It goes through the same HTTP entry point a browser
uses, then waits for the worker thread and inspects what actually landed on
disk: the plan, the tool calls, the artifacts, and the rendered bytes.

Usage:
    .venv/Scripts/python.exe web/lazy/_live_check.py "your goal here"
    .venv/Scripts/python.exe web/lazy/_live_check.py --goal "..." --timeout 900
"""

import argparse
import json
import os
import sys
import time

import django

BASE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, BASE)
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'nebians.settings')
django.setup()

from django.conf import settings  # noqa: E402
from django.test import Client  # noqa: E402

from api.models import LazyAgentRun, LazyArtifact, User  # noqa: E402

settings.ALLOWED_HOSTS = list(getattr(settings, 'ALLOWED_HOSTS', [])) + ['testserver', 'localhost']

DEFAULT_GOAL = (
    'Write a formal project report for a NEB Class 12 Computer Science '
    'investigation titled "Impact of Mobile Learning Applications on Study '
    'Habits of Secondary School Students in Nepal". It needs a title page, '
    'abstract, introduction with context on Nepali secondary education, '
    'literature review, methodology, findings and discussion, '
    'recommendations, conclusion, and references. Make it read like a real '
    'submitted academic report, not a stub. Include a diagram of the research '
    'methodology and a chart showing the distribution of study methods. '
    'Deliver it as both DOCX and PDF.'
)

FAILURES = []


def check(label, condition, detail=''):
    mark = 'PASS' if condition else 'FAIL'
    if not condition:
        FAILURES.append(label)
    print(f'  [{mark}] {label}' + (f' — {detail}' if detail else ''))


def section(title):
    print('\n' + title)
    print('-' * len(title))


def download_body(response):
    """Artifacts stream back as FileResponse, which exposes no `.content`."""
    if hasattr(response, 'streaming_content'):
        try:
            return b''.join(response.streaming_content)
        except Exception:
            return b''
    return getattr(response, 'content', b'') or b''


def pretty(value, limit=220):
    text = str(value).replace('\n', ' ')
    return text if len(text) <= limit else text[:limit] + '…'


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('goal_pos', nargs='?', default=None)
    parser.add_argument('--goal', default=None)
    parser.add_argument('--timeout', type=int, default=900)
    parser.add_argument('--model', default='neby-pro')
    args = parser.parse_args()

    goal = args.goal or args.goal_pos or DEFAULT_GOAL
    timeout = max(60, args.timeout)

    section('Setup')
    user = User.objects.exclude(id='').first()
    check('a user exists', user is not None, getattr(user, 'username', ''))
    if user is None:
        return 1

    client = Client()
    session = client.session
    session['user_id'] = user.id
    session.save()

    print(f'\nGoal: {pretty(goal, 400)}')
    print(f'Model: {args.model}   Timeout: {timeout}s')

    section('Kick off the run')
    started = time.time()
    resp = client.post(
        '/ajax/lazy/agent/run/',
        data=json.dumps({'goal': goal, 'modelKey': args.model}),
        content_type='application/json',
    )
    check('run endpoint returned 200', resp.status_code == 200, resp.status_code)
    if resp.status_code != 200:
        print(pretty(resp.content, 600))
        return 1

    payload = resp.json()
    check('run accepted', bool(payload.get('ok')), pretty(payload, 200))
    run_id = payload.get('runId')
    check('run id issued', bool(run_id), run_id)
    if not run_id:
        return 1

    print(f'runId: {run_id}')
    print(f'sessionId: {payload.get("sessionId")}')

    section('Waiting for the worker')
    terminal = {
        LazyAgentRun.STATUS_DONE, LazyAgentRun.STATUS_FAILED,
        LazyAgentRun.STATUS_CANCELLED,
    }
    last_status = None
    while time.time() - started < timeout:
        row = LazyAgentRun.objects.filter(id=run_id).values(
            'status', 'iteration', 'tool_calls', 'error',
        ).first()
        status = row['status'] if row else '?'
        if status != last_status:
            elapsed = int(time.time() - started)
            print(f'  [{elapsed:>4}s] status={status} '
                  f'iteration={row.get("iteration") if row else "?"} '
                  f'tools={row.get("tool_calls") if row else "?"}')
            last_status = status
        if status in terminal:
            break
        time.sleep(3)

    elapsed = int(time.time() - started)
    run = LazyAgentRun.objects.get(id=run_id)

    section('Outcome')
    check('run reached a terminal state', run.status in terminal, run.status)
    check('run did not fail', run.status != LazyAgentRun.STATUS_FAILED, run.error or '')
    print(f'  finished in {elapsed}s, {run.iteration} iterations, '
          f'{run.tool_calls} tool calls')

    meta = {}
    try:
        meta = json.loads(run.meta or '{}') or {}
    except Exception:
        pass
    if meta.get('summary'):
        print(f'\n  Agent summary:\n    {pretty(meta["summary"], 600)}')

    section('Plan')
    plan = {}
    try:
        plan = json.loads(run.plan_json or '{}') or {}
    except Exception:
        pass
    if plan.get('summary'):
        print(f'  {pretty(plan["summary"], 300)}')
    items = plan.get('items') or []
    check('a plan was produced', bool(items), f'{len(items)} items')
    for item in items[:15]:
        mark = {'done': '[x]', 'completed': '[x]'}.get(item.get('status'), '[ ]')
        print(f'    {mark} {pretty(item.get("text"), 110)}')

    section('Steps')
    steps = list(run.steps.all().order_by('index'))
    check('steps were persisted', bool(steps), f'{len(steps)} steps')
    tools_used = []
    for step in steps:
        detail = {}
        try:
            detail = json.loads(step.detail or '{}') or {}
        except Exception:
            pass
        if step.kind == 'tool' and detail.get('phase') == 'start':
            tools_used.append(step.title)
            print(f'    {step.index:>3}. tool  {step.title}  {pretty(detail.get("args"), 110)}')
        elif step.kind == 'think':
            print(f'    {step.index:>3}. think {pretty(step.title, 130)}')
        elif step.kind == 'plan':
            print(f'    {step.index:>3}. plan  {pretty(step.title, 130)}')
    check('at least one tool ran', bool(tools_used), ', '.join(sorted(set(tools_used))))

    # Rebuilding a document is the single most expensive mistake the agent can
    # make: it burns a full pipeline per retry and hands the user a different
    # document each time. One goal means one build.
    builds = [t for t in tools_used if t == 'build_document']
    check('the document was built once, not retried', len(builds) <= 1,
          f'{len(builds)} build_document calls')

    section('Artifacts')
    artifacts = list(LazyArtifact.objects.filter(run=run).order_by('created_at'))
    check('artifacts were produced', bool(artifacts), f'{len(artifacts)} artifacts')
    for art in artifacts:
        print(f'    {art.kind:<8} {art.name:<46} {art.size:>9,} B')

    urls = {}
    for art in artifacts:
        url = f'/ajax/lazy/agent-artifact/{run.id}/{art.path}'
        urls[art.name] = (art, url)

    section('Download check')
    for name, (art, url) in urls.items():
        r = client.get(url)
        body = download_body(r)
        ok = r.status_code == 200 and len(body) > 0
        check(f'download {name}', ok, f'{r.status_code} {len(body):,} B')

    section('Workspace files')
    r = client.get(f'/ajax/lazy/agent/run/{run.id}/files/')
    files_payload = {}
    if r.status_code == 200:
        try:
            files_payload = r.json() or {}
        except Exception:
            pass
    files = files_payload.get('files') or []
    check('workspace listing works', r.status_code == 200, r.status_code)
    check('workspace has files', bool(files), f'{len(files)} files')
    for entry in files[:40]:
        print(f'    {entry.get("path")}  ({entry.get("size", 0):,} B)')

    section('Document quality')
    # A run can hold several documents if the agent retried; the deliverable is
    # the newest one, so judge that rather than whichever came first.
    docx_art = next((a for a in reversed(artifacts) if a.name.lower().endswith('.docx')), None)
    pdf_art = next((a for a in reversed(artifacts) if a.name.lower().endswith('.pdf')), None)
    wants_pdf = 'pdf' in goal.lower()

    if docx_art:
        from web.lazy.renderers import docx_renderer  # noqa: E402
        raw = download_body(client.get(f'/ajax/lazy/agent-artifact/{run.id}/{docx_art.path}'))
        check('docx is non-trivial', len(raw) > 25000, f'{len(raw):,} B')
        try:
            import io
            document = docx_renderer.Document(io.BytesIO(raw)) if hasattr(docx_renderer, 'Document') else None
        except Exception:
            document = None
        if document is None:
            from docx import Document as _Doc
            document = _Doc(io.BytesIO(raw))
        words = sum(len(p.text.split()) for p in document.paragraphs)
        tables = len(document.tables)
        images = len(document.inline_shapes)
        headings = sum(1 for p in document.paragraphs if (p.style.name or '').lower().startswith('heading'))
        print(f'    DOCX: {words:,} words, {headings} headings, {tables} tables, {images} images')
        check('docx has substantial prose', words >= 900, f'{words} words')
        check('docx has headings', headings >= 5, f'{headings} headings')
        check('docx embeds figures', images >= 1, f'{images} inline images')

        # A document can describe a table in prose and still read well to a
        # reviewer, so the structure is checked, not just the wording.
        real_tables = [t for t in document.tables if len(t.rows) > 2 and len(t.columns) > 1]
        check('docx contains real data tables', bool(real_tables),
              f'{len(real_tables)} of {tables} tables have data')
        if real_tables:
            widest = max(real_tables, key=lambda t: len(t.rows))
            print('    largest table: '
                  f'{len(widest.rows)}x{len(widest.columns)} — '
                  + ' | '.join(c.text.strip()[:18] for c in widest.rows[0].cells))
        first = next((p.text.strip() for p in document.paragraphs if p.text.strip()), '')
        print(f'    opens with: {pretty(first, 120)}')
        print('    outline:')
        for p in document.paragraphs:
            style = (p.style.name or '').lower()
            if style.startswith('heading') and p.text.strip():
                depth = 1
                if style.startswith('heading '):
                    try:
                        depth = int(style.split(' ')[1])
                    except Exception:
                        depth = 1
                print(f'      {"  " * (depth - 1)}H{depth} {pretty(p.text.strip(), 88)}')
    else:
        check('a DOCX was produced', False, 'no docx artifact')

    if pdf_art:
        raw = download_body(client.get(f'/ajax/lazy/agent-artifact/{run.id}/{pdf_art.path}'))
        check('pdf is non-trivial', len(raw) > 20000, f'{len(raw):,} B')
        check('pdf has a valid header', raw[:5] == b'%PDF-', raw[:8])
        pages = raw.count(b'/Type /Page') or raw.count(b'/Type/Page')
        print(f'    PDF: {len(raw):,} B, ~{pages} page objects')
    elif wants_pdf:
        check('a PDF was produced', False, 'no pdf artifact')
    else:
        print('    (PDF not requested — skipping)')

    figure_files = [f for f in files if 'figures/' in (f.get('path') or '')]
    check('figures were drawn into the workspace', bool(figure_files),
          f'{len(figure_files)} figure files')

    section('History replay')
    r = client.get(f'/ajax/lazy/agent/session/{payload.get("sessionId")}/history/')
    ok = r.status_code == 200
    check('session history returns 200', ok, r.status_code)
    if ok:
        hist = r.json() or {}
        msgs = hist.get('messages') or []
        check('history contains the exchange', len(msgs) >= 2, f'{len(msgs)} messages')
        for msg in msgs:
            print(f'    {msg.get("role"):<9} {pretty(msg.get("content"), 120)}')

    section('Result')
    total = len(FAILURES)
    if total:
        print(f'{total} check(s) failed: {", ".join(FAILURES)}')
        return 1
    print('ALL LIVE CHECKS PASSED')
    return 0


if __name__ == '__main__':
    sys.exit(main())
