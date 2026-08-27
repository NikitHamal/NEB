import io
import math

from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, Emu, RGBColor

EMU_PER_PT = 12700


def _map_font(name):
    n = (name or "").lower()
    if "courier" in n:
        return "Courier New"
    if "times" in n or "georgia" in n:
        return "Times New Roman"
    if "helvetica" in n or "arial" in n or "liberation" in n:
        return "Arial"
    if "calibri" in n:
        return "Calibri"
    if "cambria" in n:
        return "Cambria"
    return None


def _size_round(s):
    return Pt(max(round(float(s) * 2) / 2, 1))


def _add_runs(paragraph, segs):
    for seg in segs:
        txt = ("\t" if seg.get("tab_before") else "") + seg["t"]
        if not txt.strip() and "\t" not in txt:
            continue
        run = paragraph.add_run(txt)
        f = run.font
        f.size = _size_round(seg["size"])
        fn = _map_font(seg.get("font", ""))
        if fn:
            f.name = fn
        col = (seg.get("color") or "#000000").lstrip("#")
        if len(col) == 6:
            try:
                f.color.rgb = RGBColor.from_string(col.upper())
            except Exception:
                pass
        if seg.get("bold"):
            f.bold = True
        if seg.get("italic"):
            f.italic = True


_NS_W = ('xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/'
         'wordprocessingDrawing"')


def _to_anchor(run_element, cx_emu, cy_emu, x_pt, y_pt, z):
    inline = run_element.find(".//" + qn("wp:inline"))
    if inline is None:
        return False
    inline.tag = qn("wp:anchor")
    for key, val in (
        ("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"),
        ("simplePos", "0"), ("relativeHeight", str(z)), ("behindDoc", "1"),
        ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1"),
    ):
        inline.set(key, val)
    extent = inline.find(qn("wp:extent"))
    effect = inline.find(qn("wp:effectExtent"))
    if extent is not None:
        extent.set("cx", str(int(cx_emu)))
        extent.set("cy", str(int(cy_emu)))
    pos_h = parse_xml(
        f'<wp:positionH {_NS_W} relativeFrom="page">'
        f"<wp:posOffset>{int(x_pt * EMU_PER_PT)}</wp:posOffset></wp:positionH>"
    )
    pos_v = parse_xml(
        f'<wp:positionV {_NS_W} relativeFrom="page">'
        f"<wp:posOffset>{int(y_pt * EMU_PER_PT)}</wp:posOffset></wp:positionV>"
    )
    wrap = parse_xml(f'<wp:wrapNone {_NS_W}/>')
    if extent is not None:
        idx = list(inline).index(extent)
        inline.insert(idx, pos_h)
        inline.insert(idx + 1, pos_v)
        anchor_at = idx + 2
        if effect is None:
            inline.insert(anchor_at, parse_xml(
                f'<wp:effectExtent {_NS_W} l="0" t="0" r="0" b="0"/>'))
            anchor_at += 1
        else:
            anchor_at += 1
        inline.insert(anchor_at, wrap)
    return True


_NS_FULL = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)


def _build_anchor_drawing(rid, cx, cy, x_pt, y_pt, z):
    xml = (
        f'<w:drawing {_NS_FULL}>'
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="{z}" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{int(x_pt * EMU_PER_PT)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{int(y_pt * EMU_PER_PT)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{int(cx)}" cy="{int(cy)}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        f'<wp:wrapNone/>'
        f'<wp:docPr id="{z}" name="lazypdf-{z}"/>'
        f'<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
        f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic>'
        f'<pic:nvPicPr><pic:cNvPr id="{z}" name="lazypdf-{z}"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr>'
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{int(cx)}" cy="{int(cy)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        f'</pic:spPr>'
        f'</pic:pic>'
        f'</a:graphicData></a:graphic>'
        f'</wp:anchor>'
        f'</w:drawing>'
    )
    return parse_xml(xml)


def _embed_image(document, paragraph, pil_img, rect, z_counter):
    if pil_img is None:
        return False
    try:
        mode = pil_img.mode
        if mode == "P":
            pil_img = pil_img.convert("RGBA")
        elif mode == "CMYK":
            pil_img = pil_img.convert("RGB")
        buf = io.BytesIO()
        pil_img.save(buf, format="PNG")
        data = buf.getvalue()
        if len(data) > 12 * 1024 * 1024:
            return False
        x_pt = rect["x"]
        y_pt = rect["y_top"]
        cx = max(int(rect["w"] * EMU_PER_PT), 9525)
        cy = max(int(rect["h"] * EMU_PER_PT), 9525)
        rid, _img = document.part.get_or_add_image(io.BytesIO(data))
        z = z_counter[0]
        drawing = _build_anchor_drawing(rid, cx, cy, x_pt, y_pt, z)
        run = paragraph.add_run()
        run._r.append(drawing)
        z_counter[0] += 1
        return True
    except Exception:
        return False


def convert_pdf_to_docx(pdf_bytes, max_pages=80, title=None):
    from docx import Document
    from . import extract as _x
    from . import plan as _p

    extracted = _x.extract_pdf(pdf_bytes)
    if not extracted["pages"]:
        raise ValueError("No pages could be read from this PDF")
    if extracted["total_runs"] == 0 and all(
        (not pg["images"]) for pg in extracted["pages"][:10]
    ):
        raise ValueError("No readable text or images found — this PDF looks like an unindexed scan")
    plan_doc = _p.build_plan(extracted, max_pages=max_pages)

    doc = Document()
    sec0 = doc.sections[0]
    doc_left = plan_doc["doc_left"]
    left_m = max(20.0, min(doc_left - 2.0, 140.0))
    needed_width = 200.0
    for pg in plan_doc["pages"]:
        for ln in pg["lines"]:
            needed_width = max(needed_width, ln["right_est"] - doc_left)

    cur_w = None
    cur_h = None
    z_counter = [251658240]
    stats = {
        "source_pages": len(plan_doc["pages"]),
        "images": 0,
        "autofit_pages": 0,
        "skipped_rotated": plan_doc.get("skipped_rotated", 0),
        "truncated": plan_doc["truncated"],
        "lines": 0,
    }
    image_budget = [20 * 1024 * 1024]

    total_pages = len(plan_doc["pages"])
    for p_i, pg in enumerate(plan_doc["pages"]):
        pw, ph = float(pg["w"]), float(pg["h"])
        if cur_w is None or abs(pw - cur_w) > 2 or abs(ph - cur_h) > 2:
            if cur_w is None:
                sec = sec0
            else:
                sec = doc.add_section(2)
            try:
                from docx.enum.section import WD_SECTION_START
                sec.start_type = WD_SECTION_START.NEW_PAGE
            except Exception:
                pass
            cur_w, cur_h = pw, ph
            sec.page_width = Emu(int(pw * EMU_PER_PT))
            sec.page_height = Emu(int(ph * EMU_PER_PT))
            rm = max(15.0, min(pw - left_m - needed_width, 45.0))
            sec.left_margin = Emu(int(left_m * EMU_PER_PT))
            sec.right_margin = Emu(int(rm * EMU_PER_PT))
            sec.top_margin = Emu(int(_p.TOP_MARGIN * EMU_PER_PT))
            sec.bottom_margin = Emu(int(_p.BOTTOM_MARGIN * EMU_PER_PT))
            new_section = True
        else:
            new_section = False

        lines = pg["lines"]
        images = sorted(pg["images"], key=lambda im: im["y_top"])
        first_para = None

        def _make_para(pbb):
            para = doc.add_paragraph()
            pf = para.paragraph_format
            if pbb:
                pf.page_break_before = True
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            return para

        if lines:
            for j, ln in enumerate(lines):
                pb_flag = (j == 0) and (not new_section) and (p_i > 0)
                para = _make_para(pb_flag)
                pf = para.paragraph_format
                pf.line_spacing = Pt(ln["lh"])
                pf.left_indent = Pt(max(ln["indent"], 0.0))
                pf.space_before = Pt(ln["space_before"])
                pf.space_after = Pt(ln["space_after"])
                if ln.get("center"):
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                seen_tabs = set()
                for tpos in ln.get("tabs_rel") or []:
                    tt = round(tpos - max(ln["indent"], 0.0), 1)
                    if tt > 14 and tt < pw - left_m - 30 and tt not in seen_tabs:
                        seen_tabs.add(tt)
                        try:
                            pf.tab_stops.add_tab_stop(Pt(tt))
                        except Exception:
                            pass
                _add_runs(para, ln["segs"])
                stats["lines"] += 1
                if j == 0:
                    first_para = para
        else:
            first_para = _make_para((not new_section) and p_i > 0)

        anchor_para = first_para or _make_para(False)
        for rect in images:
            pil = None
            if callable(pg.get("resolve_image")):
                try:
                    pil = pg["resolve_image"](rect["name"])
                except Exception:
                    pil = None
            if pil is None:
                continue
            est = (pil.width or 1) * (pil.height or 1) * 3
            if est > image_budget[0]:
                continue
            image_budget[0] -= est
            before = z_counter[0]
            if _embed_image(doc, anchor_para, pil, rect, z_counter):
                if z_counter[0] != before:
                    stats["images"] += 1

        if pg.get("autofit"):
            stats["autofit_pages"] += 1

    props = doc.core_properties
    props.title = title or plan_doc["title"] or "Converted document"
    out = io.BytesIO()
    doc.save(out)
    stats["pages_out"] = total_pages
    return {"bytes": out.getvalue(), "stats": stats}